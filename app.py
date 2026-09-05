import io
import json
import os
import re

import streamlit as st
from google import genai
from google.genai import types
from docx import Document
from pypdf import PdfReader


# -----------------------------
# Page configuration
# -----------------------------
st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload a resume to get an estimated ATS score, "
    "keyword analysis, and practical improvement suggestions."
)


# -----------------------------
# Configuration
# -----------------------------
MODEL_NAME = "gemini-2.5-flash"
MAX_FILE_SIZE_MB = 50


def get_api_key():
    """Read the Gemini API key from Streamlit secrets or environment."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None

    return key or os.getenv("GEMINI_API_KEY")


# -----------------------------
# Document extraction
# -----------------------------
def extract_pdf_text(data: bytes) -> str:
    """Extract text for deterministic checks and context."""
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def extract_docx_text(data: bytes) -> str:
    """Extract paragraphs and table content from DOCX."""
    document = Document(io.BytesIO(data))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts).strip()


# -----------------------------
# Basic deterministic checks
# -----------------------------
def basic_resume_checks(text: str) -> dict:
    """Run simple checks independent of Gemini."""
    lower = text.lower()

    sections = {
        "contact_information": any(
            term in lower
            for term in ("email", "phone", "linkedin", "github")
        ),
        "summary_or_objective": any(
            term in lower
            for term in (
                "professional summary",
                "summary",
                "objective",
                "profile",
            )
        ),
        "experience": any(
            term in lower
            for term in (
                "experience",
                "work experience",
                "employment",
            )
        ),
        "education": "education" in lower,
        "skills": "skills" in lower,
        "projects": "projects" in lower,
    }

    word_count = len(re.findall(r"\b[\w+#.-]+\b", text))
    bullet_count = len(
        re.findall(
            r"(?:^|\n)\s*(?:[-•*▪◦]|\d+[.)])\s+",
            text,
        )
    )

    return {
        "word_count": word_count,
        "bullet_count": bullet_count,
        "sections_present": sections,
        "has_email": bool(
            re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
        ),
        "has_phone": bool(
            re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", text)
        ),
    }


# -----------------------------
# Gemini analysis
# -----------------------------
def analyze_resume(file_name: str, data: bytes, job_description: str) -> dict:
    api_key = get_api_key()

    if not api_key:
        raise RuntimeError(
            "Gemini API key was not found. Add GEMINI_API_KEY "
            "to Streamlit Secrets or your environment variables."
        )

    client = genai.Client(api_key=api_key)
    lower_name = file_name.lower()

    if lower_name.endswith(".pdf"):
        mime_type = "application/pdf"
        extracted_text = extract_pdf_text(data)

        file_part = types.Part.from_bytes(
            data=data,
            mime_type=mime_type,
        )

    elif lower_name.endswith(".docx"):
        extracted_text = extract_docx_text(data)
        file_part = types.Part.from_text(text=extracted_text)

    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    checks = basic_resume_checks(extracted_text)

    prompt = """
You are an expert resume reviewer and ATS compatibility evaluator.

Analyze the uploaded resume carefully.

IMPORTANT:
The ATS score is an ESTIMATED ATS COMPATIBILITY SCORE.
It is not the score of any particular commercial ATS because
different ATS products and employers use different algorithms.

Return ONLY valid JSON using exactly this structure:

{
  "ats_score": 0,
  "score_breakdown": {
    "keyword_match": 0,
    "formatting_readability": 0,
    "section_completeness": 0,
    "achievement_quality": 0,
    "role_relevance": 0
  },
  "summary": "",
  "strengths": [],
  "improvements": [
    {
      "issue": "",
      "why_it_matters": "",
      "how_to_fix": "",
      "priority": "High"
    }
  ],
  "missing_keywords": [],
  "suggested_keywords": [],
  "section_feedback": {
    "summary": "",
    "experience": "",
    "education": "",
    "skills": "",
    "projects": ""
  },
  "ats_friendly_checklist": [
    {
      "item": "",
      "status": "Pass",
      "note": ""
    }
  ],
  "rewritten_bullets": [
    {
      "original": "",
      "improved": ""
    }
  ]
}

RULES:

1. ats_score must be an integer from 0 to 100.
2. Every score_breakdown value must be an integer from 0 to 100.
3. Never invent employers, jobs, degrees, certifications, skills,
   achievements, numbers, technologies, or qualifications.
4. If a job description is supplied, compare the resume against it.
5. Identify relevant missing keywords only when they actually appear
   important in the supplied job description.
6. If no job description is supplied, use general ATS-readiness criteria
   and state that keyword matching is generic.
7. Evaluate standard headings, keyword relevance, readability,
   achievement-focused writing, action verbs, measurable results,
   role relevance, and ATS-friendly structure.
8. Flag tables, columns, graphics, icons, unusual headings,
   headers/footers, or other ATS risks only when there is evidence.
9. Do not criticize something without evidence.
10. Rewritten bullets must improve existing resume bullets without
    inventing facts or numbers.
11. Keep the analysis concise and practical.
12. Return valid JSON only. No markdown fences.
"""

    if job_description.strip():
        prompt += (
            "\n\nTARGET JOB DESCRIPTION:\n"
            + job_description[:12000]
        )
    else:
        prompt += (
            "\n\nNo target job description was supplied. "
            "Evaluate general ATS readiness."
        )

    prompt += (
        "\n\nPROGRAMMATIC CHECKS:\n"
        + json.dumps(checks, indent=2)
    )

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=[
            file_part,
            types.Part.from_text(text=prompt),
        ],
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    try:
        result = json.loads(response.text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            "Gemini returned invalid JSON. Please try again."
        ) from exc

    return validate_result(result)


# -----------------------------
# Defensive result validation
# -----------------------------
def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(value)))
    except (TypeError, ValueError):
        return 0


def validate_result(result: dict) -> dict:
    if not isinstance(result, dict):
        raise RuntimeError("Unexpected analysis format.")

    result["ats_score"] = clamp_score(result.get("ats_score", 0))

    breakdown = result.setdefault("score_breakdown", {})
    for field in (
        "keyword_match",
        "formatting_readability",
        "section_completeness",
        "achievement_quality",
        "role_relevance",
    ):
        breakdown[field] = clamp_score(breakdown.get(field, 0))

    for field in (
        "strengths",
        "improvements",
        "missing_keywords",
        "suggested_keywords",
        "ats_friendly_checklist",
        "rewritten_bullets",
    ):
        if not isinstance(result.get(field), list):
            result[field] = []

    if not isinstance(result.get("section_feedback"), dict):
        result["section_feedback"] = {}

    if not isinstance(result.get("summary"), str):
        result["summary"] = ""

    return result


# -----------------------------
# UI helpers
# -----------------------------
def priority_icon(priority: str) -> str:
    value = priority.lower()
    if value == "high":
        return "🔴"
    if value == "medium":
        return "🟡"
    return "🟢"


# -----------------------------
# Upload area
# -----------------------------
uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="Supported formats: PDF and DOCX. Maximum size: 50 MB.",
)

job_description = st.text_area(
    "Optional: Paste the Job Description",
    height=180,
    placeholder=(
        "Paste the target job description here for "
        "job-specific keyword matching..."
    ),
)

if uploaded_file:
    st.info(
        "Tip: adding the target job description makes "
        "keyword matching much more useful."
    )

    if st.button(
        "🔍 Analyze Resume",
        type="primary",
        use_container_width=True,
    ):
        file_bytes = uploaded_file.getvalue()

        if len(file_bytes) > MAX_FILE_SIZE_MB * 1024 * 1024:
            st.error(
                f"Please upload a file smaller than "
                f"{MAX_FILE_SIZE_MB} MB."
            )
            st.stop()

        with st.spinner(
            "Analyzing your resume with Gemini Flash..."
        ):
            try:
                result = analyze_resume(
                    uploaded_file.name,
                    file_bytes,
                    job_description,
                )
                st.session_state["analysis"] = result
                st.session_state["analyzed_file"] = uploaded_file.name

            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
                st.info(
                    "Check your Gemini API key, internet connection, "
                    "file format, and try again."
                )


# -----------------------------
# Results
# -----------------------------
analysis = st.session_state.get("analysis")

if analysis:
    st.divider()

    score = analysis["ats_score"]

    col1, col2, col3 = st.columns([1, 2, 2])

    with col1:
        st.metric("Estimated ATS Score", f"{score}/100")

    with col2:
        if score >= 80:
            st.success("Strong ATS readiness")
        elif score >= 60:
            st.warning("Good start — improvements recommended")
        else:
            st.error("Needs improvement before applying")

    with col3:
        st.caption(
            "This is an AI estimate. Real ATS systems differ "
            "in parsing and ranking."
        )

    st.subheader("📊 Score Breakdown")

    breakdown = analysis["score_breakdown"]

    for label, value in breakdown.items():
        st.write(
            f"**{label.replace('_', ' ').title()} — {value}/100**"
        )
        st.progress(value / 100)

    st.subheader("📝 Overall Assessment")
    st.write(analysis["summary"])

    left, right = st.columns(2)

    with left:
        st.subheader("✅ Strengths")
        if analysis["strengths"]:
            for item in analysis["strengths"]:
                st.write(f"• {item}")
        else:
            st.write("No strengths were returned.")

    with right:
        st.subheader("🔑 Missing Keywords")
        keywords = analysis["missing_keywords"]
        if keywords:
            st.write(", ".join(str(x) for x in keywords))
        else:
            st.write(
                "No major keyword gaps were identified."
            )

    st.subheader("🛠️ Improvements")

    if analysis["improvements"]:
        for index, item in enumerate(
            analysis["improvements"],
            start=1,
        ):
            issue = item.get("issue", "Improvement")
            priority = item.get("priority", "Medium")

            with st.expander(
                f"{index}. {priority_icon(priority)} "
                f"{issue} — {priority}"
            ):
                st.write(
                    "**Why it matters:** "
                    + str(item.get("why_it_matters", ""))
                )
                st.write(
                    "**How to fix:** "
                    + str(item.get("how_to_fix", ""))
                )
    else:
        st.write("No major improvements were returned.")

    st.subheader("📌 Section Feedback")

    feedback = analysis["section_feedback"]

    for section in (
        "summary",
        "experience",
        "education",
        "skills",
        "projects",
    ):
        text = feedback.get(section, "")
        if text:
            st.markdown(
                f"**{section.replace('_', ' ').title()}**"
            )
            st.write(text)

    st.subheader("☑️ ATS-Friendly Checklist")

    checklist = analysis["ats_friendly_checklist"]

    if checklist:
        for item in checklist:
            status = str(item.get("status", "Review"))
            icon = "✅" if status.lower() == "pass" else "⚠️"

            st.write(
                f"{icon} **{item.get('item', '')}** — "
                f"{item.get('note', '')}"
            )
    else:
        st.write("No checklist items were returned.")

    if analysis["rewritten_bullets"]:
        st.subheader("✍️ Improved Bullet Examples")

        for item in analysis["rewritten_bullets"]:
            st.markdown(
                "**Original:** "
                + str(item.get("original", ""))
            )
            st.markdown(
                "**Improved:** "
                + str(item.get("improved", ""))
            )
            st.divider()

    report = json.dumps(
        analysis,
        indent=2,
        ensure_ascii=False,
    )

    st.download_button(
        "⬇️ Download Analysis JSON",
        data=report,
        file_name="resume_ats_analysis.json",
        mime="application/json",
        use_container_width=True,
    )

else:
    st.markdown(
        """
### 🚀 How it works

**1. Upload your resume**  
Upload a PDF or DOCX resume.

**2. Add a job description**  
Optional, but recommended for job-specific keyword matching.

**3. AI analysis**  
Gemini Flash analyzes ATS compatibility and resume quality.

**4. Get your results**
- 🎯 Estimated ATS score
- 📊 Score breakdown
- 🔑 Missing keywords
- ✅ Resume strengths
- 🛠️ Improvement suggestions
- ☑️ ATS-friendly checklist
- ✍️ Improved bullet examples
"""
    )
